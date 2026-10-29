"""A file an agent makes for a task is served from that task's outputs folder, previewed, and named in its brief (outputs.py).

The shape of the bug: a researcher on Codex saved a workbook under scratch/outputs/TQ-n/, cited it with a raw marker, and when
asked to send it wrote a link to /outputs/TQ-n/<file> - a route the server never had, so the owner clicked into a 404."""
import os, shutil, unittest
from pathlib import Path
from unittest import mock

from fastapi.testclient import TestClient

from taskuary import config, guard, outputs, server
from taskuary.store import MemoryStore, task_ref


def _book(path: Path):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = 'Vendors'
    ws.append(['Vendor', 'Contact', 'Annual cost'])
    ws.append(['Payworth', 'Gail Moreno', 1200])
    ws.append(['Trainly', 'Ray Colton', 480])
    ws['E1'] = '=SUM(C2:C3)'
    notes = wb.create_sheet('Notes'); notes.append(['Checked by Alex Doyle', 'x' * 900])
    path.parent.mkdir(parents=True, exist_ok=True); wb.save(path)


class TaskOutputsTests(unittest.TestCase):
    def setUp(self):
        self.store = MemoryStore()
        self.tid = self.store.create_task({'Title': 'Compare training vendors', 'Kind': 'general'}, 'owner')
        self.other = self.store.create_task({'Title': 'Another task', 'Kind': 'general'}, 'owner')
        self.dir = outputs.folder(self.tid)
        self.addCleanup(shutil.rmtree, config.home() / 'scratch' / 'outputs', True)
        _book(self.dir / 'vendor comparison.xlsx')
        (self.dir / 'notes.md').write_text('# Vendors\n\nPayworth is cheaper.', encoding='utf-8')
        (self.dir / 'page.html').write_text('<script>alert(1)</script>', encoding='utf-8')
        outputs.folder(self.other).mkdir(parents=True, exist_ok=True)
        (outputs.folder(self.other) / 'secret.txt').write_text('not yours', encoding='utf-8')
        (config.home() / 'config.toml').touch()
        p = mock.patch.object(server, 'store', self.store); p.start(); self.addCleanup(p.stop)
        self.client = TestClient(server.app)

    def test_the_folder_is_the_one_the_agent_picked(self):
        self.assertEqual(outputs.folder(17), config.home() / 'scratch' / 'outputs' / 'TQ-0017')

    def test_a_file_in_the_tasks_own_folder_is_served(self):
        listed = self.client.get(f'/api/task-outputs/{self.tid}').json()['data']
        names = {r['name']: r for r in listed}
        self.assertEqual(set(names), {'vendor comparison.xlsx', 'notes.md', 'page.html'})
        book = names['vendor comparison.xlsx']
        self.assertEqual((book['kind'], book['preview']), ('output', 'sheet'))
        self.assertNotIn(str(config.home()), str(listed))                       # a name, never where it lives on disk
        got = self.client.get(book['url'], params={'download': 'true'})
        self.assertEqual(got.status_code, 200)
        self.assertEqual(got.content, (self.dir / 'vendor comparison.xlsx').read_bytes())
        self.assertIn('attachment', got.headers['content-disposition'])
        self.assertEqual(self.client.get(f'/api/task-outputs/{self.tid}/notes.md').content, (self.dir / 'notes.md').read_bytes())

    def test_a_page_never_runs_as_taskuary(self):
        got = self.client.get(f'/api/task-outputs/{self.tid}/page.html')
        self.assertTrue(got.headers['content-type'].startswith('text/plain'))
        self.assertEqual(got.headers['x-content-type-options'], 'nosniff')

    def test_nothing_outside_the_folder_is_served(self):
        (config.home() / 'scratch' / 'outputs' / 'loose.txt').write_text('above the folder', encoding='utf-8')
        for bad in ('..%2Floose.txt', '../loose.txt', '..%5Cloose.txt', f'..%2F{task_ref(self.other)}%2Fsecret.txt',
                    'secret.txt', 'C:%2FWindows%2Fwin.ini', '%2Fetc%2Fpasswd', 'notes.md:stream', 'nothing-here.xlsx'):
            got = self.client.get(f'/api/task-outputs/{self.tid}/{bad}')       # a client may fold `..` itself: then no such route
            self.assertIn(got.status_code, (404, 422), bad)
            self.assertNotIn('above the folder', got.text); self.assertNotIn('not yours', got.text)
        for bad in ('../loose.txt', '..\\loose.txt', '/etc/passwd', 'C:/x', 'C:x', '\\\\host\\share\\x', 'a//b', './notes.md', ''):
            self.assertIsNone(outputs.resolve(self.tid, bad), bad)
        self.assertEqual(self.client.get(f'/api/task-outputs/999/notes.md').status_code, 404)

    def test_a_link_out_of_the_folder_is_not_followed(self):
        outside = config.home() / 'outside.txt'; outside.write_text('mine', encoding='utf-8')
        try: os.symlink(outside, self.dir / 'link.txt')
        except (OSError, NotImplementedError): self.skipTest('this machine cannot make a symlink')
        self.assertIsNone(outputs.resolve(self.tid, 'link.txt'))
        self.assertNotIn('link.txt', {r['name'] for r in outputs.listing(self.tid)})

    def test_a_workbook_previews_as_rows(self):
        got = self.client.get(f'/api/task-outputs/{self.tid}/vendor%20comparison.xlsx', params={'preview': 'true'}).json()
        self.assertEqual(got['sheet'], 'Vendors')
        self.assertEqual(got['sheets'], ['Vendors', 'Notes'])
        self.assertEqual(got['rows'][0][:3], ['Vendor', 'Contact', 'Annual cost'])
        self.assertEqual(got['rows'][1][:3], ['Payworth', 'Gail Moreno', '1200'])
        self.assertFalse(got['more_rows'])
        notes = self.client.get(f'/api/task-outputs/{self.tid}/vendor%20comparison.xlsx', params={'preview': 'true', 'sheet': 'Notes'}).json()
        self.assertEqual(notes['sheet'], 'Notes')
        self.assertLessEqual(len(notes['rows'][0][1]), outputs.CELL)               # a cell is capped
        self.assertEqual(self.client.get(f'/api/task-outputs/{self.tid}/notes.md', params={'preview': 'true'}).status_code, 422)

    def test_rows_and_columns_are_capped(self):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        for i in range(outputs.ROWS + 50): ws.append([f'r{i}c{j}' for j in range(outputs.COLS + 10)])
        wb.save(self.dir / 'big.xlsx')
        got = outputs.sheet(self.dir / 'big.xlsx')
        self.assertEqual((len(got['rows']), len(got['rows'][0])), (outputs.ROWS, outputs.COLS))
        self.assertTrue(got['more_rows'] and got['more_cols'])

    def test_without_openpyxl_the_workbook_is_download_only(self):
        with mock.patch.object(outputs, 'can_sheet', lambda: False):
            row = next(r for r in outputs.listing(self.tid) if r['name'].endswith('.xlsx'))
            self.assertIsNone(row['preview'])
            self.assertIn('Download', row['note'])
            got = self.client.get(f'/api/task-outputs/{self.tid}/vendor%20comparison.xlsx', params={'preview': 'true'})
            self.assertEqual(got.status_code, 501)
            self.assertEqual(self.client.get(row['url'], params={'download': 'true'}).status_code, 200)

    def test_the_assistant_payload_lists_them(self):
        got = self.client.get(f'/api/tasks/{self.tid}/assistant').json()
        self.assertIn('vendor comparison.xlsx', {r['name'] for r in got['outputs']})

    def test_a_download_link_carries_the_token_in_its_query(self):
        "An <a href> cannot send a header: the outputs route takes the token in the query, as attachments do."
        cfg = {**server.cfg, 'server': {**server.cfg['server'], 'token': 'owner-secret'}}
        with mock.patch.object(server, 'cfg', cfg):
            url = f'/api/task-outputs/{self.tid}/notes.md'
            self.assertEqual(self.client.get(url).status_code, 401)
            self.assertEqual(self.client.get(url, params={'token': 'owner-secret'}).status_code, 200)

    def test_a_strangers_session_reaches_only_its_own_tasks_outputs(self):
        d = lambda p: guard.stranger_denied('GET', p, {}, self.tid, lambda k, a: None)
        self.assertEqual(d(f'/api/task-outputs/{self.tid}'), '')
        self.assertEqual(d(f'/api/task-outputs/{self.tid}/notes.md'), '')
        self.assertTrue(d(f'/api/task-outputs/{self.other}/secret.txt'))
        self.assertTrue(d(f'/api/task-outputs/{self.tid}9/x'))

    def test_a_cli_chat_is_told_where_its_files_go(self):
        from taskuary import general
        sess = general.GeneralSession(self.store, self.tid)
        sess.pick, sess.provider, sess.model = 'cli:codex', 'codex', ''
        seen = []
        with mock.patch.object(general.llm_mod, 'build_llm', return_value=lambda system, user, **kw: seen.append(system) or 'Done.'):
            sess.send_prompt('make the comparison a spreadsheet')
        self.assertIn(str(outputs.folder(self.tid)), seen[0])
        self.assertIn('Never write a URL', seen[0])

    def test_the_brief_names_the_folder_and_forbids_made_up_addresses(self):
        said = outputs.brief(self.tid)
        self.assertIn(str(outputs.folder(self.tid)), said)
        self.assertIn('Never write a URL', said)
        self.assertTrue(outputs.folder(self.tid).is_dir())


def _stamp(p: Path, n: int):
    "Files written n minutes apart, so the newest is known."
    t = 1_790_000_000 + n * 60
    os.utime(p, (t, t))


class OneOfEachTests(unittest.TestCase):
    """The owner, 2026-10-28: "it looks like it created a million attachments". A workbook saved v1..v9 beside its helper script, its log
    and a temp folder is ONE deliverable: its newest version, the eight before it behind it, the working files apart."""
    def setUp(self):
        self.store = MemoryStore()
        self.tid = self.store.create_task({'Title': 'Vendor pricing review', 'Summary': 'Compare what Payworth and Trainly charge.',
                                           'Kind': 'general'}, 'owner')
        self.other = self.store.create_task({'Title': 'Snag list for the new office', 'Kind': 'general'}, 'owner')
        self.addCleanup(shutil.rmtree, config.home() / 'scratch' / 'outputs', True)
        d = outputs.folder(self.tid); d.mkdir(parents=True, exist_ok=True)
        for i in range(1, 10):
            (d / f'pricing comparison v{i}.xlsx').write_bytes(b'x' * i); _stamp(d / f'pricing comparison v{i}.xlsx', i)
        for name in ('build_sheet.py', 'run.log'): (d / name).write_text('scratch', encoding='utf-8'); _stamp(d / name, 20)
        (d / 'tmp').mkdir(); (d / 'tmp' / 'rates.csv').write_text('a,b', encoding='utf-8'); _stamp(d / 'tmp' / 'rates.csv', 21)
        (d / 'summary.pdf').write_bytes(b'%PDF'); _stamp(d / 'summary.pdf', 5)
        o = outputs.folder(self.other); o.mkdir(parents=True, exist_ok=True)
        (o / 'snag list.xlsx').write_bytes(b'x'); (o / 'plan.md').write_text('# Plan', encoding='utf-8')

    def test_nine_versions_and_scratch_are_one_file_and_eight_earlier(self):
        made, working = outputs.deliverables(outputs.listing(self.tid))
        self.assertEqual([r['name'] for r in made], ['pricing comparison v9.xlsx', 'summary.pdf'])
        self.assertEqual(len(made[0]['earlier']), 8)
        self.assertEqual(made[0]['earlier'][0]['name'], 'pricing comparison v8.xlsx')
        self.assertEqual(sorted(r['path'] for r in working), ['build_sheet.py', 'run.log', 'tmp/rates.csv'])

    def test_a_family_is_one_name_without_its_version_marks(self):
        one = {outputs.family(n) for n in ('Vendor pricing v9.xlsx', 'vendor_pricing_final.xlsx', 'vendor pricing (2).xlsx',
                                           'vendor-pricing-20261027-1530.xlsx', 'Vendor Pricing_v2_final.xlsx', 'drafts/vendor pricing draft.xlsx')}
        self.assertEqual(one, {'vendor pricing.xlsx'})
        # ...and different things stay different: another year, another kind, a set the ask wanted
        self.assertEqual(len({outputs.family(n) for n in ('budget 2025.xlsx', 'budget 2026.xlsx', 'budget.pdf', 'region 1.csv', 'region 2.csv')}), 5)
        self.assertEqual(outputs.family('renew.docx'), 'renew.docx')

    def test_the_phone_offers_each_thing_once(self):
        from taskuary import remote_assistant
        self.assertEqual([r['name'] for r in remote_assistant.task_files(self.store, {'tid': self.tid})], ['pricing comparison v9.xlsx', 'summary.pdf'])

    def test_every_worker_hears_one_final_file(self):
        from taskuary import terminal
        self.assertIn('one final file', outputs.brief(self.tid))
        self.assertIn('one final file', self.store.doc('agent') or open(Path(outputs.__file__).parent / 'templates' / 'agent.md', encoding='utf-8').read())
        older = terminal.with_fence_rule('Rules from before files had a rule.')
        self.assertIn('one final file', older)
        self.assertIn('scratch work', older)
        self.assertEqual(terminal.with_fence_rule(older), older)

    def test_the_assistant_finds_the_excel_on_the_task_about_vendor_pricing(self):
        from taskuary import concierge
        # what a file is about is read by meaning (embeddings, textsim), never by counting shared word stems (2026-10-28)
        from tests import fake_embed
        meaning = fake_embed.meaning({'pricing': ['pricing'], 'snag': ['snag']})
        meaning.__enter__(); self.addCleanup(meaning.__exit__, None, None, None)
        hits = outputs.find(self.store, 'vendor pricing', 'xlsx')
        self.assertEqual([(h['tid'], h['name']) for h in hits], [(self.tid, 'pricing comparison v9.xlsx')])
        said = concierge.read_op(self.store, 'outputs.find', {'query': 'the excel on the task about vendor pricing', 'kind': 'xlsx'})
        self.assertIn('pricing comparison v9.xlsx', said)
        self.assertIn(f'/api/task-outputs/{self.tid}/pricing%20comparison%20v9.xlsx', said)
        self.assertIn(f'#task={self.tid}', said)
        self.assertIn('8 earlier versions', said)
        self.assertNotIn('build_sheet.py', said)
        # by the file's own name, by one task, and a miss says so
        self.assertIn('snag list.xlsx', concierge.read_op(self.store, 'outputs.find', {'query': 'snag', 'kind': 'excel'}))
        self.assertIn('plan.md', concierge.read_op(self.store, 'outputs.find', {'ref': task_ref(self.other)}))
        self.assertNotIn('pricing', concierge.read_op(self.store, 'outputs.find', {'ref': task_ref(self.other)}))
        self.assertIn('No file', concierge.read_op(self.store, 'outputs.find', {'query': 'payroll accruals', 'kind': 'pdf'}))

    def test_on_a_phone_no_link_is_given(self):
        from taskuary import concierge
        tok = concierge.DELIVERY.set(concierge.PHONE)
        try: said = concierge.read_op(self.store, 'outputs.find', {'query': 'vendor pricing'})
        finally: concierge.DELIVERY.reset(tok)
        self.assertIn('pricing comparison v9.xlsx', said)
        self.assertNotIn('/api/task-outputs', said)

    def test_the_catalogue_offers_it_within_the_index(self):
        from taskuary import toolcatalog
        b = toolcatalog.block()
        self.assertIn('"outputs.find"', b)
        self.assertLess(len(b), 7950)
        self.assertEqual(toolcatalog.valid('outputs.find', {}), '')
        self.assertTrue(toolcatalog.is_read('outputs.find'))


if __name__ == '__main__':
    unittest.main()
