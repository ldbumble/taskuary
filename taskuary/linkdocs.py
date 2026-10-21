"""Documents a message POINTS AT, as first-class as the files it carries: one registry, one reader per source.

"Here are all the outstanding bills - what do you want to pay?" arrives as a LINK to a shared sheet, and the link was a
word in the body: triage judged it blind, the drafter could not quote a row, and an agent asked to "go through them and
compare with last month" had nothing to go through. A link to a Google Doc/Sheet/Slides/Drive file, a OneDrive or SharePoint
file, a GitHub issue or pull request, a file on the web (a contract PDF) or a plain web page is now a DOCUMENT on the message:

- REGISTRY. `find(text)` turns a body into documents - {id, url, kind, label, source, ref, name}. Derived from the stored
  body on every read (no table, no migration, every old message has its documents too); the id is the URL's hash, so the
  same sheet quoted under four replies is one document on the task (`of_task`).
- RESOLVERS. One per source, `RESOLVERS[source](store, doc, full, offset, limit)`, each reaching the document through the card
  that owns it - google: the Google Sheets card's token (Sheets API for a sheet, Drive for the rest); graph: the SharePoint
  card's app or the Outlook card's tenant app (/shares); github: the GitHub card; web: a plain fetch behind webguard. A new
  source is one function and one line in RESOLVERS. A document no card can reach says why in one line, with the card that
  fixes it: `Outstanding bills (Google Sheet) (not read: no Google sign-in with file access - ...)`.
- STRUCTURE STAYS. A sheet or a csv/xlsx reaches a reader as ROWS - csv lines, header first, ROWS of them, with the total
  count and how to get the rest - never flattened prose, so "compare with previous months" has columns to compare.
- EVERY READER, ONE ROAD. `entries()` hands attachtext the same {name, text, chars} / {name, note} entries a file makes, so
  triage, the drafter, the Advisor and the agent's context file read linked documents exactly as they read attachments:
  bounded, scrubbed and fenced as untrusted. An agent re-reads one in full by id (server: /api/tasks/<id>/documents).

Bounds: MAX_DOCS documents per message; a plain web page counts only while a message has at most MAX_PAGES of them (a
newsletter's thirty links are navigation, not documents); a prompt sees ROWS rows; a full read returns at most FULL_ROWS rows
or FULL_CHARS characters at a time. Reads are cached (attachtext's cache) for TTL, a failure for NOTE_TTL, so the poll pays
once and a card fixed in the meantime is tried again soon.

WHO MAY MAKE US OPEN A LINK. A credentialed read opens the link with the OWNER's accounts. A stranger who sends a link to
the owner's own payroll sheet would have it read into a draft addressed back to them - so a document from a sender the owner
has never written to is listed, not opened, at intake (senders.known); the owner, or an agent the owner starts on it, can
read it on request. A plain web page is fetched only on request too: fetching every link at intake tells a tracker the mail
was read. A file on the web (.pdf, .xlsx ...) is the document itself and is fetched, through webguard (no private addresses).
"""
import base64, csv, hashlib, io, re, time
from urllib.parse import parse_qs, quote, unquote, urlsplit
from loguru import logger

VERSION = 1
MAX_DOCS, MAX_PAGES = 8, 3
ROWS, FULL_ROWS, FULL_CHARS = 60, 500, 60000
TTL, NOTE_TTL = 6 * 3600, 600
FILE_EXTS = ('pdf', 'docx', 'xlsx', 'pptx', 'csv', 'tsv', 'txt', 'md', 'json')
TABLE_EXTS = ('csv', 'tsv', 'xlsx')
_SKIP_EXTS = ('png', 'jpg', 'jpeg', 'gif', 'svg', 'webp', 'ico', 'css', 'js')
_URL = re.compile(r'https?://[^\s<>"\'`]+', re.I)
SHEETS, DRIVE = 'https://sheets.googleapis.com/v4/spreadsheets', 'https://www.googleapis.com/drive/v3'
GRAPH, GH = 'https://graph.microsoft.com/v1.0', 'https://api.github.com'
G_APPS = 'application/vnd.google-apps.'
# where each source is fixed: the card that holds the credentials (words.json places are for pages, these are cards)
FIX = {'google': {'label': 'Open Google Sheets', 'href': '#connector=google_sheets'},
       'graph': {'label': 'Open SharePoint', 'href': '#connector=sharepoint'},
       'github': {'label': 'Open GitHub', 'href': '#connector=github'}}


class Unreadable(Exception):
    "A document a resolver could not read - `note` is the one line a reader is shown, `fix` the card that changes it."
    def __init__(self, note, fix=None): super().__init__(note); self.note, self.fix = note, fix


# ── the registry: what a URL is ───────────────────────────────────────────────────────────────
def _unwrap(url: str) -> str:
    "A Safe Links or google.com/url wrapper around the real address, unwrapped - the wrapper is the mail system's, not the link."
    p = urlsplit(url)
    host, q = (p.hostname or '').lower(), parse_qs(p.query)
    if host.endswith('safelinks.protection.outlook.com') and q.get('url'): return q['url'][0]
    if host in ('www.google.com', 'google.com') and p.path == '/url' and (q.get('q') or q.get('url')): return (q.get('q') or q['url'])[0]
    return url


def _ext(path: str) -> str:
    last = unquote(path.rstrip('/').rsplit('/', 1)[-1])
    return last.rsplit('.', 1)[-1].lower() if '.' in last else ''


def _fname(path: str) -> str: return unquote(path.rstrip('/').rsplit('/', 1)[-1])[:120]


def classify(url: str) -> dict | None:
    "{id, url, kind, label, source, ref, name} for a link worth calling a document - None for mail chrome, a picture, a mailto."
    url = _unwrap(str(url or '').strip().rstrip('.,;:!?)]>\'"'))
    try: p = urlsplit(url)
    except ValueError: return None
    host, path = (p.hostname or '').lower(), p.path or ''
    if p.scheme not in ('http', 'https') or not host: return None
    def doc(kind, label, source, ref='', name=''):
        return {'id': 'D' + hashlib.sha1(url.encode()).hexdigest()[:7], 'url': url, 'kind': kind, 'label': label,
                'source': source, 'ref': ref, 'name': name or label}
    if host == 'docs.google.com' and (m := re.match(r'/(spreadsheets|document|presentation|forms)/d/(?:e/)?([\w-]{10,})', path)):
        kind, label = {'spreadsheets': ('sheet', 'Google Sheet'), 'document': ('doc', 'Google Doc'),
                       'presentation': ('slides', 'Google Slides'), 'forms': ('form', 'Google Form')}[m.group(1)]
        gid = (re.search(r'gid=(\d+)', p.fragment + '&' + p.query) or [None, ''])[1]
        return {**doc(kind, label, 'google', m.group(2)), **({'tab': gid} if gid else {})}
    if host == 'drive.google.com':
        if m := re.match(r'/file/d/([\w-]{10,})', path): return doc('file', 'Google Drive file', 'google', m.group(1))
        if m := re.match(r'/drive/(?:u/\d+/)?folders/([\w-]{10,})', path): return doc('folder', 'Google Drive folder', 'google', m.group(1))
        if (fid := parse_qs(p.query).get('id')): return doc('file', 'Google Drive file', 'google', fid[0])
        return None
    if host.endswith('.sharepoint.com') or host in ('onedrive.live.com', '1drv.ms'):
        one = host in ('onedrive.live.com', '1drv.ms') or '-my.sharepoint.com' in host
        code = (re.match(r'/:(\w):/', path) or [None, ''])[1]
        kind = {'x': 'sheet', 'w': 'doc', 'p': 'slides', 'b': 'pdf', 'f': 'folder'}.get(code) or \
               ('sheet' if _ext(path) in TABLE_EXTS else 'pdf' if _ext(path) == 'pdf' else 'file')
        if not code and not _ext(path) and host.endswith('.sharepoint.com') and '/sites/' not in path and '/personal/' not in path: return None
        word = {'sheet': 'spreadsheet', 'doc': 'document', 'slides': 'presentation', 'pdf': 'PDF', 'folder': 'folder'}.get(kind, 'file')
        return doc(kind, f"{'OneDrive' if one else 'SharePoint'} {word}", 'graph', name=_fname(path) if _ext(path) else '')
    if host in ('www.dropbox.com', 'dropbox.com') and re.match(r'/(s|scl|sh)/', path): return doc('file', 'Dropbox file', None, name=_fname(path))
    if (host == 'box.com' or host.endswith('.box.com')) and re.match(r'/s/', path): return doc('file', 'Box file', None)
    if host == 'github.com' and (m := re.match(r'/([\w.-]+/[\w.-]+)/(issues|pull)/(\d+)', path)):
        kind = 'issue' if m.group(2) == 'issues' else 'pull'
        return doc(kind, 'GitHub issue' if kind == 'issue' else 'GitHub pull request', 'github', f'{m.group(1)}#{m.group(3)}',
                   f'{m.group(1)}#{m.group(3)}')
    ext = _ext(path)
    if ext in _SKIP_EXTS: return None
    if ext in FILE_EXTS: return doc('table' if ext in TABLE_EXTS else ext, ext.upper() if ext != 'pdf' else 'PDF', 'web', name=_fname(path))
    return doc('page', 'Web page', 'web', name=f"{host}{path if path not in ('', '/') else ''}"[:120])


def is_doc_link(url: str) -> bool:
    "Whether an <a href> in HTML mail is a document worth keeping in the body's text (channels._clean) - a known host or a file."
    d = classify(url)
    return bool(d) and d['kind'] != 'page'


def find(text: str, skip=()) -> list:
    "The documents a body points at, first mention first, at most MAX_DOCS - plain pages only while there are MAX_PAGES or fewer."
    out, seen = [], set(skip or ())
    for u in _URL.findall(str(text or '')):
        d = classify(u)
        if d and d['url'] not in seen and d['id'] not in {x['id'] for x in out}: seen.add(d['url']); out.append(d)
    if sum(d['kind'] == 'page' for d in out) > MAX_PAGES: out = [d for d in out if d['kind'] != 'page']
    return out[:MAX_DOCS]


def _row_msg(m: dict) -> dict:
    "An ingest dict or a stored message row, as one shape: the words, who wrote them, and where."
    g = lambda a, b: m.get(a) if m.get(a) is not None else m.get(b)
    return {'body': g('body', 'BodyText') or '', 'channel': g('channel', 'Channel') or '', 'from_email': g('from_email', 'FromEmail'),
            'from_name': g('from_name', 'FromName'), 'source_name': g('source_name', 'SourceName'), 'source_link': g('source_link', 'SourceLink'),
            'conversation_id': g('conversation_id', 'ConversationId'), 'direction': g('direction', 'Direction'),
            'status': g('status', 'Status'), 'MailMetaJson': m.get('MailMetaJson'), 'mid': m.get('MessageId')}


def of_msg(m: dict) -> list:
    """The documents on one message. Taskuary's own reports and posts carry none (they are not a sender pointing at a file),
    and a GitHub item's links to other issues are the tracker's cross-references, not documents sent to the owner."""
    from .store import OWN_CHANNELS
    m = _row_msg(m)
    if m['channel'] in OWN_CHANNELS: return []
    got = find(m['body'], skip=[m['source_link']] if m['source_link'] else ())
    return [d for d in got if not (m['channel'] == 'github' and d['source'] == 'github')]


def of_message(store, mid) -> list: return [{**d, 'mid': mid} for d in of_msg(store.get_message(mid) or {})]


def of_task(store, tid) -> list:
    "Every document on a task's messages, once each - the one quoted under four replies is one document - oldest mention first."
    out = {}
    for m in store.list_messages(int(tid)) or []:
        for d in of_msg(m): out.setdefault(d['id'], {**d, 'mid': m['MessageId']})
    # ...and the task's own description: a hand-off from the chat ("go through the bills in <link>") has no message, only the
    # owner's words - which are the owner's, so mid None reads as trusted. Only a task the OWNER made: triage writes a mailed
    # task's description from the sender's words, and a stranger's link must not borrow the owner's trust through it
    t = store.get_task(int(tid)) or {}
    if t.get('OriginBase') == 'owner':
        for d in find(t.get('Summary') or ''): out.setdefault(d['id'], {**d, 'mid': None})
    return list(out.values())


def find_on_task(store, tid, did) -> dict | None:
    "A document on this task by its id - the only documents a full read reaches (never an arbitrary URL a caller names)."
    return next((d for d in of_task(store, tid) if d['id'] == did or d['url'] == did), None)


# ── reading: the cache, the gates, the resolvers ─────────────────────────────────────────────
def _key(d): return f"link:{hashlib.sha256(d['url'].encode()).hexdigest()}-v{VERSION}"


def peek(d: dict) -> dict | None:
    "What the last read of this document found, from the cache only - no network. The page and the router read this."
    from . import attachtext
    hit = attachtext._cached(_key(d))
    return hit.get('got') if isinstance(hit, dict) else None


def _trusted(store, m: dict) -> bool:
    "May an intake read open this message's links with the owner's accounts? The owner's own words and senders.known say yes."
    m = _row_msg(m)
    if m['direction'] == 'out' or m['status'] == 'context': return True
    try:
        from . import senders
        return bool(senders.known(store, m, exclude_mid=m['mid'])[0])
    except Exception as e: logger.debug(f'linkdocs: sender check failed - {e}'); return False


def read(store, d: dict, *, auto: bool = True, trusted: bool = True, fresh: bool = False, full: bool = False,
         offset: int = 0, limit: int = 0) -> dict:
    """{'text', 'chars', title?, rows?, total_rows?, more?} or {'note', fix?} for one document - never raises. `auto` is an intake
    or prompt read (the gates apply, the cache answers); a full read (`full`) is an agent's or the owner's and pages through."""
    from . import attachtext
    src = d.get('source')
    if not src: return {'note': f"no connected card reads {d['label']} links yet - open it yourself"}
    if auto and d['kind'] == 'page': return {'note': 'a web page - read on request (an agent can fetch it by its id)'}
    if auto and not trusted: return {'note': 'sent by someone you have not written to, so it is not opened with your accounts - '
                                              'open it yourself, or start an agent on it'}
    key = _key(d)
    if auto and not fresh and not full and isinstance(hit := attachtext._cached(key), dict):
        if time.time() - float(hit.get('at') or 0) < (NOTE_TTL if 'note' in (hit.get('got') or {}) else TTL): return hit['got']
    try: got = RESOLVERS[src](store, d, full, max(0, int(offset or 0)), int(limit or 0) or (FULL_ROWS if full else ROWS))
    except Unreadable as e: got = {'note': e.note, **({'fix': e.fix} if e.fix else {})}
    except Exception as e:
        logger.debug(f"linkdocs: {d['url'][:80]} unreadable - {e}")
        from .redact import scrub
        got = {'note': f'could not be read ({scrub(str(e))[:120]})'}
    if 'text' in got:
        from .redact import scrub
        text = got['text'] if 'rows' in got or not full else got['text'][offset:]     # a table pages by rows, a text by characters
        got = {'chars': len(got['text']), **got, 'text': scrub(text)[:FULL_CHARS if full else attachtext.KEEP_CHARS]}
        if full and 'rows' not in got: got['offset'] = offset
        if 'rows' in got: got['rows'] = [[scrub('' if v is None else str(v)) for v in r] for r in got['rows']]
    if not full: attachtext._remember(key, {'got': got, 'at': time.time()})
    return got


def name_of(d: dict, got: dict = None) -> str:
    "`Outstanding bills (Google Sheet)` - the document's own title when a read found one, else its file name, else its kind."
    t = str((got or {}).get('title') or d.get('name') or '').strip()
    return f"{t} ({d['label']})" if t and t != d['label'] else d['label']


def entries(store, m: dict, auto: bool = True) -> list:
    """attachtext entries for the documents a message (an ingest dict or a stored row) points at - the same {name, text, chars}
    or {name, note} a file makes, plus `doc` (the id, so a document quoted on three messages is shown once) and `more`."""
    docs = of_msg(m)
    if not docs: return []
    ok = _trusted(store, m)
    out = []
    for d in docs:
        got = read(store, d, auto=auto, trusted=ok)
        out.append({'name': name_of(d, got), 'doc': d['id'], 'link': True, 'url': d['url'],
                    **{k: got[k] for k in ('text', 'chars', 'note', 'more') if k in got}})
    return out


# ── tables: rows stay rows ───────────────────────────────────────────────────────────────────
def _csv(rows) -> str:
    b = io.StringIO()
    csv.writer(b, lineterminator='\n').writerows([['' if v is None else str(v) for v in r] for r in rows])
    return b.getvalue().strip()


def table(title: str, rows: list, offset: int, limit: int, tab: str = '', tabs=(), **more) -> dict:
    """A table as a reader gets it: the header row, then `limit` rows from `offset`, as csv - with the total and what is left.
    `rows` is a list of lists, the header first (a sheet's values, a csv, an xlsx tab)."""
    rows = [list(r) for r in rows or [] if any(str(v or '').strip() for v in r)]
    head, body = (rows[0], rows[1:]) if rows else ([], [])
    part = body[offset:offset + limit]
    text = _csv([head] + part) if head else ''
    left = len(body) - offset - len(part)
    where = f' in tab "{tab}"' + (f' (tabs: {", ".join(tabs)})' if len(tabs) > 1 else '') if tab else ''
    note = (f'rows {offset + 1}-{offset + len(part)} of {len(body)}{where}; {left} more - read the rest by the document\'s id'
            if left > 0 or offset else f'all {len(body)} rows{where}')
    return {'title': title, 'text': text or '(empty)', 'chars': len(text), 'rows': [head] + part, 'total_rows': len(body),
            'offset': offset, 'tab': tab, 'tabs': list(tabs), 'more': note, **more}


def _file(name: str, data: bytes, offset: int, limit: int, **more) -> dict:
    "Bytes of a file a link pointed at: a table as rows, anything else through attachtext's extractors."
    from . import attachtext
    ext = _ext(name)
    if ext in TABLE_EXTS:
        try:
            if ext == 'xlsx':
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                ws, tabs = wb.active, wb.sheetnames
                rows = [list(r) for _, r in zip(range(20000), ws.iter_rows(values_only=True))]
                t = table(name, rows, offset, limit, ws.title, tabs, **more); wb.close(); return t
            text = data.decode('utf-8-sig', errors='replace')
            delim = '\t' if ext == 'tsv' else (csv.Sniffer().sniff(text[:4000], delimiters=',;\t|').delimiter if text.strip() else ',')
            return table(name, list(csv.reader(io.StringIO(text), delimiter=delim)), offset, limit, **more)
        except Exception as e: logger.debug(f'linkdocs: {name} not read as a table - {e}')
    got = attachtext.read(name, data)
    if 'note' in got: raise Unreadable(got['note'])
    return {'title': name, 'text': got['text'], 'chars': got['chars'], **more}


# ── the resolvers, one per source ───────────────────────────────────────────────────────────
def _http(url, tok=None, params=None, headers=None, raw=False, who='the site'):
    "One GET for every resolver (tests replace this). Raises Unreadable on a status a reader should hear about in words."
    from . import net
    r = net.get(url, headers={**({'Authorization': f'Bearer {tok}'} if tok else {}), **(headers or {})}, params=params, timeout=20)
    if r.status_code >= 300: raise _refused(r, who)
    return r.content if raw else r.json()


def _refused(r, who) -> Unreadable:
    body = str(getattr(r, 'text', '') or '')[:400]
    if who == 'google':
        if r.status_code == 401: return Unreadable("Google refused the Google Sheets card's sign-in - sign in again", FIX['google'])
        if r.status_code == 403 and re.search(r'insufficient|scope', body, re.I):
            return Unreadable("no Google sign-in with file access - the Google Sheets card's token needs drive.readonly", FIX['google'])
        if r.status_code in (403, 404): return Unreadable("not shared with the Google account on the Google Sheets card")
    if who == 'graph':
        if r.status_code in (401, 403): return Unreadable('the Microsoft app has no file access - it needs Files.Read.All or Sites.Read.All', FIX['graph'])
        if r.status_code == 404: return Unreadable('not found, or not shared with the Microsoft app')
    if who == 'github':
        if r.status_code in (401, 403): return Unreadable('the GitHub card cannot read that repository', FIX['github'])
        if r.status_code == 404: return Unreadable('not found, or the GitHub card cannot see that repository')
    return Unreadable(f'{who} answered {r.status_code}')


def _google(store, d, full, offset, limit):
    from . import sheets
    cfg = sheets.google_sheets_connection(store)
    if not (cfg.get('google_refresh_token') and cfg.get('google_client_id')):
        raise Unreadable('no Google sign-in with file access - connect the Google Sheets card', FIX['google'])
    try: tok = sheets._token(cfg)
    except RuntimeError as e: raise Unreadable(f'the Google sign-in did not work ({str(e)[:100]})', FIX['google'])
    g = lambda url, **p: _http(url, tok, p, who='google')
    if d['kind'] == 'sheet':
        got = _gsheet(g, d['ref'], d.get('tab'), offset, limit)
        if not full: return got
        try: return {**got, 'related': _gsiblings(g, g(f"{DRIVE}/files/{d['ref']}", fields='id,parents', supportsAllDrives='true'))}
        except Unreadable as e: return {**got, 'related': [], 'related_note': e.note}      # Sheets scope alone cannot list a folder
    meta = g(f"{DRIVE}/files/{d['ref']}", fields='id,name,mimeType,size,parents', supportsAllDrives='true')
    mt, name = str(meta.get('mimeType') or ''), meta.get('name') or d['label']
    rel = {'related': _gsiblings(g, meta)} if full else {}
    if mt == G_APPS + 'spreadsheet': return {**_gsheet(g, d['ref'], None, offset, limit), **rel}
    if mt == G_APPS + 'folder':
        kids = g(f'{DRIVE}/files', q=f"'{d['ref']}' in parents and trashed=false", orderBy='modifiedTime desc', pageSize=100,
                 fields='files(id,name,mimeType,modifiedTime,webViewLink)', supportsAllDrives='true', includeItemsFromAllDrives='true')
        return table(name, [['name', 'kind', 'modified', 'link']] + [[f.get('name'), f.get('mimeType', '').replace(G_APPS, ''),
                     str(f.get('modifiedTime') or '')[:10], f.get('webViewLink')] for f in kids.get('files') or []], offset, limit)
    from . import attachtext
    if mt.startswith(G_APPS):
        text = _http(f"{DRIVE}/files/{d['ref']}/export", tok, {'mimeType': 'text/plain'}, raw=True, who='google').decode('utf-8', 'replace')
        return {'title': name, 'text': text.strip(), 'chars': len(text.strip()), **rel}
    if int(meta.get('size') or 0) > attachtext.READ_BYTES: raise Unreadable(f"{attachtext._mb(int(meta['size']))}, over the read limit")
    return _file(name, _http(f"{DRIVE}/files/{d['ref']}", tok, {'alt': 'media', 'supportsAllDrives': 'true'}, raw=True, who='google'),
                 offset, limit, **rel)


def _gsheet(g, sid, gid, offset, limit):
    meta = g(f'{SHEETS}/{sid}', fields='properties.title,sheets.properties(sheetId,title)')
    tabs = [s.get('properties') or {} for s in meta.get('sheets') or []]
    tab = next((t for t in tabs if gid and str(t.get('sheetId')) == str(gid)), tabs[0] if tabs else {'title': 'Sheet1'})
    vals = g(f"{SHEETS}/{sid}/values/{quote(tab['title'], safe='')}", valueRenderOption='FORMATTED_VALUE').get('values') or []
    return table((meta.get('properties') or {}).get('title') or 'Google Sheet', vals, offset, limit, tab['title'], [t.get('title') for t in tabs])


def _gsiblings(g, meta) -> list:
    "The other files in the same Drive folder, newest first - last month's sheet usually sits beside this one's."
    try:
        par = (meta.get('parents') or [None])[0]
        if not par: return []
        kids = g(f'{DRIVE}/files', q=f"'{par}' in parents and trashed=false", orderBy='modifiedTime desc', pageSize=25,
                 fields='files(id,name,mimeType,modifiedTime,webViewLink)', supportsAllDrives='true', includeItemsFromAllDrives='true')
        return [{'name': f.get('name'), 'modified': str(f.get('modifiedTime') or '')[:10], 'url': f.get('webViewLink')}
                for f in kids.get('files') or [] if f.get('id') != meta.get('id')][:20]
    except Unreadable: return []


def _graph(store, d, full, offset, limit):
    from . import sharepoint, attachtext
    cfg = sharepoint.sharepoint_connection(store)
    if not (cfg.get('client_id') and cfg.get('client_secret')):
        raise Unreadable('no Microsoft app with file access - set up the SharePoint card, or the Outlook card as a tenant app', FIX['graph'])
    try: tok = sharepoint._token(cfg)
    except RuntimeError as e: raise Unreadable(f'the Microsoft app did not sign in ({str(e)[:100]})', FIX['graph'])
    share = 'u!' + base64.urlsafe_b64encode(d['url'].encode()).decode().rstrip('=')
    g = lambda path, raw=False, **p: _http(f'{GRAPH}{path}', tok, p or None, raw=raw, who='graph')
    it = g(f'/shares/{share}/driveItem')
    name, par = it.get('name') or d.get('name') or d['label'], it.get('parentReference') or {}
    rel = {}
    if full and par.get('driveId') and par.get('id'):
        try:
            kids = g(f"/drives/{par['driveId']}/items/{par['id']}/children", **{'$top': 50})
            rel = {'related': [{'name': k.get('name'), 'modified': str(k.get('lastModifiedDateTime') or '')[:10], 'url': k.get('webUrl')}
                               for k in sorted(kids.get('value') or [], key=lambda k: str(k.get('lastModifiedDateTime') or ''), reverse=True)
                               if k.get('id') != it.get('id')][:20]}
        except Unreadable: pass
    if it.get('folder'):
        kids = g(f'/shares/{share}/driveItem/children', **{'$top': 200})
        return table(name, [['name', 'kind', 'modified', 'link']] + [[k.get('name'), 'folder' if k.get('folder') else 'file',
                     str(k.get('lastModifiedDateTime') or '')[:10], k.get('webUrl')] for k in kids.get('value') or []], offset, limit)
    if int(it.get('size') or 0) > attachtext.READ_BYTES: raise Unreadable(f"{attachtext._mb(int(it['size']))}, over the read limit")
    return _file(name, g(f'/shares/{share}/driveItem/content', raw=True), offset, limit, **rel)


def _github(store, d, full, offset, limit):
    c = store.get_connector_by_type('github', with_secret=True) or {}
    if not (c.get('Active') and c.get('Secret')): raise Unreadable('no GitHub card connected', FIX['github'])
    repo, n = d['ref'].split('#')
    h = {'Accept': 'application/vnd.github+json', 'X-GitHub-Api-Version': '2022-11-28'}
    i = _http(f'{GH}/repos/{repo}/issues/{n}', c['Secret'], headers=h, who='github')
    cs = _http(f'{GH}/repos/{repo}/issues/{n}/comments', c['Secret'], {'per_page': 100 if full else 20}, headers=h, who='github') if i.get('comments') else []
    who = lambda x: (x.get('user') or {}).get('login') or '?'
    text = '\n\n'.join([f"{i.get('title') or ''} [{i.get('state')}{', a pull request' if i.get('pull_request') else ''}] by {who(i)}",
                        str(i.get('body') or '').strip()] + [f"{who(x)} ({str(x.get('created_at') or '')[:10]}): {str(x.get('body') or '').strip()}" for x in cs])
    return {'title': f"{repo}#{n} {i.get('title') or ''}".strip(), 'text': text, 'chars': len(text)}


def _web(store, d, full, offset, limit):
    from . import webguard, attachtext
    try: r = webguard.get(d['url'], timeout=20, headers={'User-Agent': 'Taskuary (reading a link sent to its owner)'})
    except RuntimeError as e: raise Unreadable(str(e)[:160])
    if r.status_code >= 300: raise Unreadable(f'the site answered {r.status_code}')
    data, ct = r.content or b'', str(r.headers.get('content-type') or '').lower()
    if len(data) > attachtext.READ_BYTES: raise Unreadable(f'{attachtext._mb(len(data))}, over the read limit')
    if 'html' in ct or (d['kind'] == 'page' and not ct):
        from .channels import _clean
        html = data.decode(r.encoding or 'utf-8', errors='replace')
        title = (re.search(r'(?is)<title[^>]*>(.*?)</title>', html) or [None, ''])[1].strip()[:120]
        text = _clean(html)
        return {'title': title or d['name'], 'text': text, 'chars': len(text)}
    return _file(d.get('name') or 'file', data, offset, limit)


RESOLVERS = {'google': _google, 'graph': _graph, 'github': _github, 'web': _web}
