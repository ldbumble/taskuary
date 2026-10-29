"""The files an agent makes FOR the owner on a task - a spreadsheet, a PDF, a chart - and how the task page reaches them.

A chat's CLI works in Taskuary's scratch folder (llm._build_llm), and nothing told it where a finished file should go. A Codex
researcher saved a workbook under scratch/outputs/TQ-0017/ by its own habit, cited it with its own `:codex-file-citation{...}`
marker - printed raw in the chat - and, asked to send it, made up http://127.0.0.1:7787/outputs/TQ-0017/... : a route this server
never had. Session artifacts (session_artifacts.py) are only files Taskuary itself copied in under ~/.taskuary/artifacts, so a
file the agent wrote was on disk and unreachable from the page.

Now the folder is named: <home>/scratch/outputs/TQ-nnnn/ for task nnnn - the very one the agent picked - it is said in the brief
(brief), and GET /api/task-outputs/<tid>/<file> serves a file in THAT task's folder and nothing else (resolve)."""
import mimetypes, re
from datetime import datetime
from pathlib import Path

from . import config
from .store import task_ref

MAX_FILES, MAX_DEPTH = 200, 4
ROWS, COLS, CELL, SHEETS = 200, 40, 300, 30          # what a preview carries: enough to recognise the sheet, never the whole book
TYPES = {'.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.xlsm': 'application/vnd.ms-excel.sheet.macroEnabled.12',
         '.csv': 'text/csv', '.tsv': 'text/tab-separated-values', '.md': 'text/markdown', '.txt': 'text/plain', '.json': 'application/json',
         '.log': 'text/plain', '.pdf': 'application/pdf', '.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg',
         '.gif': 'image/gif', '.webp': 'image/webp', '.svg': 'image/svg+xml', '.html': 'text/html', '.htm': 'text/html',
         '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation'}
# how the page previews each kind; anything not here is Download only
PREVIEW = {'.xlsx': 'sheet', '.xlsm': 'sheet', '.csv': 'table', '.pdf': 'pdf', '.png': 'picture', '.jpg': 'picture', '.jpeg': 'picture',
           '.gif': 'picture', '.webp': 'picture', '.md': 'text', '.txt': 'text', '.json': 'text', '.log': 'text', '.tsv': 'text',
           '.svg': 'page', '.html': 'page', '.htm': 'page', '.docx': 'document'}


def folder(tid: int) -> Path: return config.home() / 'scratch' / 'outputs' / task_ref(tid)
def content_type(name: str) -> str: return TYPES.get(Path(name).suffix.lower()) or mimetypes.guess_type(name)[0] or 'application/octet-stream'


def can_sheet() -> bool:
    "openpyxl ships in the desktop exe (deps.BUNDLE); a plain pip install may not have it - then a workbook is Download only."
    try: import openpyxl  # noqa: F401
    except ImportError: return False
    return True


def _base_ok(base: Path) -> bool:
    "The folder itself is ours too: a scratch/outputs that is a link to somewhere else serves nothing."
    return not any(q.is_symlink() for q in (base.parent, base)) and base.resolve().is_relative_to((config.home() / 'scratch').resolve())


_DRIVE = re.compile(r'^[A-Za-z]:|^[\\/]{2}')

def resolve(tid: int, rel: str):
    """The file `rel` names inside task `tid`'s outputs folder, or None. A name only: never absolute, never a drive or a share, no
    `..` or empty step, no `:` (an NTFS stream), no link anywhere on the way - and after resolving, still under the folder."""
    rel = str(rel or '')
    if not rel or '\x00' in rel or _DRIVE.match(rel) or rel[0] in '/\\': return None
    parts = rel.replace('\\', '/').split('/')
    if any(p in ('', '.', '..') or ':' in p for p in parts): return None
    base = folder(tid)
    try:
        if not _base_ok(base): return None
        p, cur = base.joinpath(*parts), base
        for part in parts:
            cur = cur / part
            if cur.is_symlink(): return None
        p = p.resolve()
        return p if p.is_relative_to(base.resolve()) and p.is_file() else None
    except (OSError, ValueError):
        return None


def row(tid: int, p: Path, rel: str) -> dict:
    "One file as the page draws it - the same shape as a shown page (server._artifact_row), so one list draws both."
    from urllib.parse import quote
    ext, sheet = p.suffix.lower(), PREVIEW.get(p.suffix.lower())
    if sheet == 'sheet' and not can_sheet(): sheet = None
    st = p.stat()
    return {'id': f'output:{rel}', 'name': p.name, 'path': rel, 'kind': 'output', 'content_type': content_type(p.name),
            'size': st.st_size, 'at': datetime.fromtimestamp(st.st_mtime).strftime('%Y-%m-%d %H:%M:%S'), 'preview': sheet, 'url': f'/api/task-outputs/{int(tid)}/{quote(rel)}',
            **({'note': 'No preview on this install - Download opens it in your spreadsheet app.'} if PREVIEW.get(ext) == 'sheet' and not sheet else {})}


def listing(tid: int) -> list:
    "Every file in the task's outputs folder, newest first. Lock files and hidden files are not outputs; links are never followed."
    base = folder(tid)
    try:
        if not base.is_dir() or not _base_ok(base): return []
    except (OSError, ValueError): return []
    out, stack = [], [(base, 0)]
    while stack and len(out) < MAX_FILES:
        d, depth = stack.pop()
        try: kids = sorted(d.iterdir())
        except OSError: continue
        for k in kids:
            if k.name.startswith(('.', '~$')) or k.is_symlink(): continue
            if k.is_dir() and depth < MAX_DEPTH: stack.append((k, depth + 1))
            elif k.is_file():
                rel = k.relative_to(base).as_posix()
                if resolve(tid, rel): out.append((k.stat().st_mtime, row(tid, k, rel)))
    return [r for _, r in sorted(out, key=lambda x: -x[0])][:MAX_FILES]


# ONE OF EACH THING IT MADE (the owner, 2026-10-28: "it looks like it created a million attachments. Agents should not create lists of
# lists of attachments unless specifically asked"). A workbook saved as v1..v9 listed nine rows, with the helper script and its log
# between them. A FAMILY is one name with its version marks taken off (v2, _final, (1), a timestamp); only its newest is the deliverable,
# the rest are its earlier versions, kept and reachable behind it. The agent's own working files - its scripts, logs, temp and cache
# files, anything under a tmp/ or scratch/ folder - are not things it made for the owner, so they fold under one line.
# taskOutputs.deliverables is the same rule for the page: keep the two alike.
_VER = re.compile(r'(?:[ _.-]+(?:v|ver|version|rev)[ _.-]?\d+|[ _.-]+(?:final|draft|copy|old|new|latest|updated|edited|revised)'
                  r'|[ _.-]*\(\d+\)|[ _.-]+\d{8}(?:[ _T-]?\d{4,6})?|[ _.-]+\d{4}-\d\d-\d\d(?:[ _T-]?\d\d[-.]?\d\d(?:[-.]?\d\d)?)?)+$', re.I)
_WORK_EXT = {'.log', '.tmp', '.temp', '.bak', '.pyc', '.cache', '.lock', '.swp', '.part', '.crdownload',
             '.py', '.ps1', '.sh', '.bat', '.cmd', '.js', '.mjs', '.ts', '.sql', '.r', '.ipynb'}
_WORK_DIR = re.compile(r'(?:^|/)(?:tmp|temp|scratch|cache|__pycache__|work|working|intermediate|drafts?)/', re.I)
_WORK_NAME = re.compile(r'^(?:tmp|temp|scratch|debug)(?:[ _.-]|$)', re.I)


def family(name: str) -> str:
    "A file's name without its version marks: 'Vendor pricing v9.xlsx', 'vendor pricing (2).xlsx' and 'vendor_pricing_final.xlsx' are one."
    p = Path(str(name or '').replace('\\', '/').rsplit('/', 1)[-1].lower())
    stem = _VER.sub('', p.stem).strip(' _.-') or p.stem
    return re.sub(r'[ _.-]+', ' ', stem) + p.suffix


def is_working(rel: str) -> bool:
    "The agent's own working file - a script, a log, a temp or cache file, anything in a tmp/ or scratch/ folder - not a deliverable."
    rel = str(rel or '').replace('\\', '/')
    return Path(rel).suffix.lower() in _WORK_EXT or bool(_WORK_DIR.search(rel)) or bool(_WORK_NAME.match(rel.rsplit('/', 1)[-1]))


def deliverables(rows: list) -> tuple:
    """(made, working): each family once - its newest row, the older ones on it as `earlier`, newest first - and the working files apart.
    A row that is not a file in the folder (a page the agent showed or published) is never a working file: it chose to show it."""
    made, by, working = [], {}, []
    for r in sorted(rows or [], key=lambda r: str(r.get('at') or r.get('created_at') or ''), reverse=True):
        if r.get('kind') == 'output' and is_working(r.get('path') or r.get('name')): working.append(r); continue
        k = family(r.get('name') or r.get('path'))
        if k in by: by[k].setdefault('earlier', []).append(r)
        else: by[k] = dict(r); made.append(by[k])
    return made, working


# WHERE IS THE EXCEL THE AGENT MADE (the owner, 2026-10-28: "if I ask the chat assistant to pull up the Excel created by the agent on the
# task about X, could it find the link to it? It should be able to"). Every task's outputs folder and the pages its agents showed or
# published, matched by the task's words and the file's own name - the model passes the words and the kind; nothing reads the sentence.
_SHEET = ('.xlsx', '.xlsm', '.xls', '.csv', '.tsv')
KINDS = {'xlsx': _SHEET, 'sheet': _SHEET, 'spreadsheet': _SHEET, 'workbook': _SHEET, 'excel': _SHEET, 'csv': ('.csv', '.tsv'),
         'pdf': ('.pdf',), 'doc': ('.docx', '.doc', '.md', '.txt', '.rtf'), 'document': ('.docx', '.doc', '.md', '.txt', '.rtf', '.pdf'),
         'docx': ('.docx', '.doc'), 'word': ('.docx', '.doc'), 'deck': ('.pptx', '.ppt'), 'pptx': ('.pptx', '.ppt'), 'slides': ('.pptx', '.ppt'),
         'picture': ('.png', '.jpg', '.jpeg', '.gif', '.webp', '.svg'), 'chart': ('.png', '.jpg', '.jpeg', '.svg', '.html'),
         'page': ('.html', '.htm', '.md'), 'text': ('.md', '.txt')}


def exts_of(kind: str) -> tuple:
    "The extensions a `kind` means: a name in KINDS or an extension itself ('xlsx', '.pdf'); none means any file."
    k = str(kind or '').strip().lower().lstrip('.')
    return KINDS.get(k) or KINDS.get(k.rstrip('s')) or ((f'.{k}',) if re.fullmatch(r'[a-z0-9]{1,5}', k) else ())


def _tids() -> list:
    base = config.home() / 'scratch' / 'outputs'
    try: return sorted({int(m.group(1)) for d in base.iterdir() if d.is_dir() and (m := re.fullmatch(r'TQ-0*(\d+)', d.name))}, reverse=True)
    except OSError: return []


def _page(a: dict) -> dict:
    "A page an agent showed or published, in the row shape (server._artifact_row) - here, so a look-up never imports the server."
    return {'id': a['ArtifactId'], 'name': a.get('Name') or 'page', 'kind': a.get('Kind'), 'content_type': a.get('ContentType') or '',
            'created_at': a.get('CreatedAt'), 'url': f"/api/task-artifacts/{a['ArtifactId']}" if a.get('Path') else None,
            **({'external_url': a['Url']} if a.get('Url') else {})}


def find(store, query: str = '', kind: str = '', tid: int = None, limit: int = 8) -> list:
    """The files agents made, best match first: [{tid, title, name, url, at, earlier, ...}] - each family once (deliverables), working
    files never. `query` is matched against the task's title and summary and the file's own name by MEANING (embeddings, textsim) -
    a count of shared word stems decided it once (the owner, 2026-10-28) - or, with no embedder, found there whole; `kind` narrows
    by what it is. Newest first when nothing is asked about."""
    from . import textsim
    want, q = exts_of(kind), ' '.join(str(query or '').split())
    ids, pages = [int(tid)] if tid else _tids(), {}
    try:
        for a in store._rows("SELECT * FROM task_artifact WHERE Kind IN ('shown','claude_artifact')" + (' AND TaskId=?' if tid else ''),
                             (int(tid),) if tid else ()):
            pages.setdefault(int(a['TaskId']), []).append(_page(a))
    except Exception: pass
    hits = []
    for t in dict.fromkeys(ids + ([] if tid else sorted(pages, reverse=True))):
        task = store.get_task(t)
        if not task: continue
        made, _ = deliverables(listing(t) + pages.get(t, []))
        about = f"{task.get('Title') or ''} {task.get('Summary') or ''}"
        for r in made:
            if want and Path(str(r.get('name') or '')).suffix.lower() not in want: continue
            text = f"{re.sub(r'[_.-]+', ' ', str(r.get('name') or ''))} - {about}"
            hits.append((f"file:{t}:{r.get('name')}", text, str(r.get('at') or r.get('created_at') or ''), {'tid': t, 'title': task.get('Title') or '', **r}))
    if not q: return [h[3] for h in sorted(hits, key=lambda h: h[2], reverse=True)][:limit]
    near = textsim.rank(store, q, [(k, text) for k, text, *_ in hits])
    by = {h[0]: h for h in hits}
    if near is not None: return [by[k][3] for k, _ in near][:limit]
    return [h[3] for h in sorted(hits, key=lambda h: h[2], reverse=True) if q.lower() in ' '.join(h[1].split()).lower()][:limit]


def _cell(v) -> str:
    if v is None: return ''
    s = v.isoformat(sep=' ') if hasattr(v, 'isoformat') else str(v)
    return s if len(s) <= CELL else s[:CELL - 1] + '…'


def sheet(p: Path, name: str = '') -> dict:
    """A workbook's sheet as rows the page draws: the named one (default the first), capped at ROWS x COLS, each cell at CELL
    characters; values, not formulas (data_only), read without loading the whole book (read_only)."""
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        ws = wb[name] if name in names else wb[names[0]]
        rows, more = [], False
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i >= ROWS: more = True; break
            rows.append([_cell(v) for v in r[:COLS]])
        while rows and not any(rows[-1]): rows.pop()
        width = max((max((j + 1 for j, c in enumerate(r) if c), default=0) for r in rows), default=0)
        cut = (ws.max_column or 0) > COLS
        return {'sheet': ws.title, 'sheets': names[:SHEETS], 'rows': [r[:width] + [''] * (width - len(r[:width])) for r in rows],
                'more_rows': more, 'more_cols': cut, 'max_rows': ROWS, 'max_cols': COLS}
    finally: wb.close()


DOC_BLOCKS = 400                                     # a document preview: its first paragraphs and tables, never the whole book
_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def document(p: Path) -> dict:
    """A Word document as the page draws it: its paragraphs (a heading marked as one) and its tables as rows, in order, the first
    DOC_BLOCKS of them - read from the file's own XML (a .docx is a zip), so no library is needed and nothing in it runs."""
    import zipfile
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(p) as z: root = ET.fromstring(z.read('word/document.xml'))
    body = root.find(f'{_W}body')
    text = lambda el: ''.join(t.text or '' for t in el.iter(f'{_W}t'))
    blocks, more = [], False
    for el in list(body) if body is not None else []:
        if len(blocks) >= DOC_BLOCKS: more = True; break
        if el.tag == f'{_W}p':
            style = el.find(f'{_W}pPr/{_W}pStyle')
            head = style is not None and str(style.get(f'{_W}val') or '').lower().startswith(('heading', 'title'))
            if (t := text(el).strip()) or blocks: blocks.append({'text': _cell(t), **({'heading': True} if head else {})})
        elif el.tag == f'{_W}tbl':
            rows = [[_cell(text(c).strip()) for c in r.findall(f'{_W}tc')][:COLS] for r in el.findall(f'{_W}tr')][:ROWS]
            if rows: blocks.append({'rows': rows})
    while blocks and not blocks[-1].get('text', True): blocks.pop()
    return {'blocks': blocks, 'more': more, 'max_blocks': DOC_BLOCKS}


def brief(tid: int) -> str:
    "What a chat with hands is told about the files it makes - so it saves them where the page finds them and never invents an address."
    d = folder(tid)
    try: d.mkdir(parents=True, exist_ok=True)
    except OSError: pass
    return (f'FILES YOU MAKE FOR THE OWNER (a spreadsheet, a document, a PDF, a chart): save each one in {d} . Every file there shows on '
            'this task with Preview and Download, so the owner opens it from the chat. To point at one, write its full path in your '
            'answer. Never write a URL or a link to it - Taskuary serves that folder itself, and an address you make up leads nowhere. '
            f'Asked to send it "here", it already is: name the file. {ONE_FILE}')


# ...and a TERMINAL session is told the same by its environment (terminal.session_env: TASKUARY_OUTPUTS) and its standing rules
# (AGENT.md, or WHERE appended to an older copy - terminal.with_fence_rule), not its seed, which a real TUI types character by
# character under a budget: a Claude Code session on a task saved every workbook it made to the owner's Downloads and a scratchpad,
# where the task page - serving only this folder - had nothing to open or download.
ENV = 'TASKUARY_OUTPUTS'
WHERE = (f"Save each file you make for the owner in the task's outputs folder - its path is in ${ENV} (or your brief): the task page "
         'offers only that folder to open and download, so a file saved anywhere else never reaches the owner.')


def env_folder(tid) -> str:
    "The folder a session on task `tid` saves into, made now so the agent finds it there; '' with no task."
    if not str(tid or '').isdigit(): return ''
    d = folder(int(tid))
    try: d.mkdir(parents=True, exist_ok=True)
    except OSError: pass
    return str(d)


# ...and at the source (the owner, 2026-10-28): an agent that leaves every draft and every helper beside the answer is what made the list
ONE_FILE = ('Deliver one final file per thing asked: overwrite it as you improve it (or keep your drafts elsewhere), never leave every '
            'version beside it. Keep your scratch work - scripts, logs, test data, intermediate files - out of the folder the owner reads '
            'from. Make several files, or a list of files, only when the ask wants several.')
